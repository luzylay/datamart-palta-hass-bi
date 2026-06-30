import docx
import sys

def find_placeholders(filename):
    try:
        doc = docx.Document(filename)
        for i, p in enumerate(doc.paragraphs):
            text = p.text.lower()
            if "captura" in text or "imagen" in text or "plot" in text or "gráfico" in text or "grafico" in text or "nota:" in text:
                print(f"Para {i}: {p.text}")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    find_placeholders(sys.argv[1])
