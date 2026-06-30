import docx
import sys
import re

def parse_md(filename):
    # simple parser for Todas_medida_sin_filtrar.md
    with open(filename, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    measures = {}
    for line in lines:
        if line.startswith('|') and not line.startswith('| MEASURE') and not line.startswith('| ---'):
            parts = [p.strip() for p in line.split('|')]
            if len(parts) >= 5:
                name = parts[1]
                table = parts[2]
                formula = parts[3].replace('\\n', '\n').replace('\\_', '_').replace('\\*', '*')
                measures[name] = formula
    return measures

def update_docx(docx_in, docx_out, measures):
    doc = docx.Document(docx_in)
    
    updates = 0
    
    # Update DAX formulas in 1x1 tables
    for i, table in enumerate(doc.tables):
        if len(table.rows) == 1 and len(table.columns) == 1:
            cell = table.cell(0, 0)
            text = cell.text.strip()
            if " = " in text:
                name_in_doc = text.split(" = ")[0].strip()
                # Try to find a match in measures
                for m_name, m_formula in measures.items():
                    if m_name in name_in_doc or name_in_doc in m_name:
                        # Found a match!
                        new_text = f"{m_name} = \n{m_formula}"
                        if text != new_text:
                            cell.text = new_text
                            updates += 1
                            print(f"Updated DAX for {m_name}")
                        break
                        
    doc.save(docx_out)
    print(f"Total updates: {updates}")

if __name__ == "__main__":
    measures = parse_md('info-para-actualizar-usar/Todas_medida_sin_filtrar.md')
    print(f"Found {len(measures)} measures in MD.")
    update_docx('Proyecto_IB_Grupo01.docx', 'Proyecto_IB_Grupo01_updated.docx', measures)
